"""
Generate formio-pitch-deck.pptx from 11 HTML pitch deck slides.
Recreates the visual structure using python-pptx with matching colors, fonts, and layout.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu, Cm
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from copy import deepcopy

# Constants
BG_COLOR = RGBColor(0xFF, 0xFF, 0xEB)  # #ffffeb warm cream
NAVY = RGBColor(0x0A, 0x13, 0x22)       # #0a1322 deep navy
BLUE = RGBColor(0x00, 0x88, 0xFF)       # #0088ff electric blue
MUTED = RGBColor(0x5A, 0x5A, 0x52)      # #5a5a52
LIGHT_MUTED = RGBColor(0x9A, 0x9A, 0x8E)  # #9a9a8e
GHOST = RGBColor(0xF0, 0xF0, 0xE4)      # very faint for ghost numbers
GRAY_AA = RGBColor(0xAA, 0xAA, 0xAA)    # #aaaaaa

HEADING_FONT = 'Georgia'
BODY_FONT = 'Calibri'
BRAND_FONT = 'Calibri'

# Slide dimensions: 16:9
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Margins (roughly 12vw = ~1.6in from each side, 8vh = ~0.6in top/bottom)
LEFT_MARGIN = Inches(1.6)
RIGHT_MARGIN = Inches(1.6)
TOP_MARGIN = Inches(0.75)
CONTENT_WIDTH = SLIDE_WIDTH - LEFT_MARGIN - RIGHT_MARGIN


def set_slide_bg(slide, color):
    """Set solid background color for a slide."""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height):
    """Add a textbox and return the shape."""
    return slide.shapes.add_textbox(left, top, width, height)


def set_text(tf, text, font_name=BODY_FONT, size=Pt(14), color=NAVY, bold=False, italic=False, alignment=PP_ALIGN.LEFT):
    """Set text in a text frame with formatting."""
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic


def add_run(paragraph, text, font_name=BODY_FONT, size=Pt(14), color=NAVY, bold=False, italic=False):
    """Add a run to an existing paragraph."""
    run = paragraph.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = size
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    return run


def add_slide_number(slide, num_str):
    """Add slide number in bottom-right."""
    left = SLIDE_WIDTH - Inches(2.0)
    top = SLIDE_HEIGHT - Inches(0.6)
    txBox = add_textbox(slide, left, top, Inches(1.0), Inches(0.4))
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    add_run(p, num_str, BODY_FONT, Pt(11), LIGHT_MUTED, bold=False)


def create_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Use blank layout
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: Title slide
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Logo / Wordmark "Formio" top-left
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(0.5), Inches(3), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Formio", BRAND_FONT, Pt(36), NAVY, bold=True)

    # Three statements stacked vertically, centered vertically
    statements = [
        ("Collecte l'info client.", NAVY, False),
        ("Genere les documents legaux.", NAVY, False),
        ("Automatiquement.", BLUE, True),
    ]
    y_start = Inches(2.5)
    gap = Inches(1.2)
    for i, (text, color, is_italic) in enumerate(statements):
        txBox = add_textbox(slide, LEFT_MARGIN, y_start + gap * i, Inches(9), Inches(1.0))
        tf = txBox.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        add_run(p, text, HEADING_FONT, Pt(44), color, bold=False, italic=is_italic)

    # Founder name bottom-right
    txBox = add_textbox(slide, SLIDE_WIDTH - Inches(4.5), SLIDE_HEIGHT - Inches(0.9), Inches(3.5), Inches(0.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    add_run(p, "Parsa Homayouni, Fondateur", BODY_FONT, Pt(11), LIGHT_MUTED)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.RIGHT
    add_run(p2, "Equipe 2 -- ENT-4020", BODY_FONT, Pt(9), LIGHT_MUTED)

    # =========================================================================
    # SLIDE 2: Problem
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Left column: title
    left_w = Inches(5.5)
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(2.0), left_w, Inches(2.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "Les cabinets ", HEADING_FONT, Pt(38), NAVY)
    add_run(p, "perdent", HEADING_FONT, Pt(38), BLUE, italic=True)
    add_run(p, " des heures par dossier.", HEADING_FONT, Pt(38), NAVY)

    # Subtitle
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(4.5), left_w, Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Leurs clients remplissent mal les formulaires.", BODY_FONT, Pt(16), MUTED)

    # Right column: 3 numbered points
    right_x = Inches(7.5)
    right_w = Inches(5.0)
    points = [
        ("1.", "Allers-retours avec le client"),
        ("2.", "Temps non-facturable perdu"),
        ("3.", "Risque d'erreur legale"),
    ]
    y = Inches(2.0)
    for num, text in points:
        txBox = add_textbox(slide, right_x, y, right_w, Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        add_run(p, num + "  ", HEADING_FONT, Pt(26), BLUE)
        add_run(p, text, BODY_FONT, Pt(22), NAVY, bold=False)
        y += Inches(1.5)

    add_slide_number(slide, "02")

    # =========================================================================
    # SLIDE 3: Insight - 20%
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Label
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.2), Inches(3), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "L'INSIGHT", BODY_FONT, Pt(11), LIGHT_MUTED, bold=True)

    # Huge stat "20%"
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.8), Inches(10), Inches(3.0))
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    add_run(p, "20", HEADING_FONT, Pt(144), NAVY, bold=True)
    add_run(p, "%", HEADING_FONT, Pt(144), BLUE, bold=True, italic=True)

    # Description
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(5.0), Inches(8), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "C'est tout ce qu'on demande au client.", HEADING_FONT, Pt(22), NAVY)

    # Detail
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(5.8), Inches(9), Inches(0.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "Le reste est extrait des documents (OCR) ou reduit a des yes/no triviaux.", BODY_FONT, Pt(14), MUTED)

    add_slide_number(slide, "03")

    # =========================================================================
    # SLIDE 4: How it works (centered)
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Eyebrow centered
    txBox = add_textbox(slide, Inches(0), Inches(0.8), SLIDE_WIDTH, Inches(0.6))
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "COMMENT CA MARCHE", BODY_FONT, Pt(16), LIGHT_MUTED, bold=True)

    # Headline centered
    txBox = add_textbox(slide, Inches(0), Inches(1.5), SLIDE_WIDTH, Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "Deux etapes, ", HEADING_FONT, Pt(60), GRAY_AA, italic=True)
    add_run(p, "c'est tout", HEADING_FONT, Pt(60), BLUE, italic=True)

    # Two columns: Step 01 and Step 02
    col1_x = Inches(1.5)
    col2_x = Inches(7.5)
    col_w = Inches(5.0)
    step_y = Inches(3.5)

    # Ghost "01"
    txBox = add_textbox(slide, col1_x, step_y - Inches(0.3), col_w, Inches(2.0))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "01", HEADING_FONT, Pt(96), GHOST)

    # Step 1 title
    txBox = add_textbox(slide, col1_x + Inches(0.3), step_y + Inches(1.2), col_w, Inches(0.7))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Le questionnaire client", HEADING_FONT, Pt(24), NAVY)

    # Step 1 desc
    txBox = add_textbox(slide, col1_x + Inches(0.3), step_y + Inches(1.9), Inches(4.5), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "Votre client repond a des questions simples depuis son telephone.", BODY_FONT, Pt(13), MUTED)

    # Connector text centered
    txBox = add_textbox(slide, Inches(5.0), step_y + Inches(0.8), Inches(3.5), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "Formio s'occupe du reste  -->", BODY_FONT, Pt(11), LIGHT_MUTED, italic=True)

    # Ghost "02"
    txBox = add_textbox(slide, col2_x, step_y - Inches(0.3), col_w, Inches(2.0))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "02", HEADING_FONT, Pt(96), GHOST)

    # Step 2 title
    txBox = add_textbox(slide, col2_x + Inches(0.3), step_y + Inches(1.2), col_w, Inches(0.7))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Les documents generes", HEADING_FONT, Pt(24), NAVY)

    # Step 2 desc
    txBox = add_textbox(slide, col2_x + Inches(0.3), step_y + Inches(1.9), Inches(4.5), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "FORMIO genere les formulaires IMM, Arrima et IRCC automatiquement.", BODY_FONT, Pt(13), MUTED)

    add_slide_number(slide, "04")

    # =========================================================================
    # SLIDE 5: Traction
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Label
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(0.8), Inches(3), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "TRACTION", BODY_FONT, Pt(11), LIGHT_MUTED, bold=True)

    # Title
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.3), Inches(10), Inches(1.2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "2 mois. 4 cabinets. Revenus ", HEADING_FONT, Pt(40), NAVY)
    add_run(p, "cette semaine.", HEADING_FONT, Pt(40), BLUE, italic=True)

    # Subtitle
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(2.5), Inches(8), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Zero marketing. Zero outreach. Tout organique.", BODY_FONT, Pt(14), MUTED, bold=False)

    # Left column: stats
    stat_y = Inches(3.5)
    # ~10K$
    txBox = add_textbox(slide, LEFT_MARGIN, stat_y, Inches(4), Inches(1.2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "~10K$", HEADING_FONT, Pt(54), NAVY, bold=True)
    p2 = tf.add_paragraph()
    add_run(p2, "ARR signe (2 cabinets payants)", BODY_FONT, Pt(12), MUTED)

    # 4
    txBox = add_textbox(slide, LEFT_MARGIN, stat_y + Inches(1.8), Inches(4), Inches(1.2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "4", HEADING_FONT, Pt(54), NAVY, bold=True)
    p2 = tf.add_paragraph()
    add_run(p2, "cabinets en onboarding imminent", BODY_FONT, Pt(12), MUTED)

    # Right column: partners
    partner_x = Inches(7.0)
    # PARTENAIRES label
    txBox = add_textbox(slide, partner_x, Inches(3.1), Inches(5), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "PARTENAIRES", BODY_FONT, Pt(10), LIGHT_MUTED, bold=True)

    # AQAADI
    txBox = add_textbox(slide, partner_x, Inches(3.6), Inches(5.5), Inches(1.2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "AQAADI", HEADING_FONT, Pt(22), NAVY, bold=True)
    p2 = tf.add_paragraph()
    add_run(p2, "via Me Gagner, ambassadrice technologique", BODY_FONT, Pt(12), MUTED)
    p3 = tf.add_paragraph()
    add_run(p3, "(acces aux 500+ membres)", BODY_FONT, Pt(12), MUTED)

    # Blouin Avocats
    txBox = add_textbox(slide, partner_x, Inches(5.2), Inches(5.5), Inches(1.2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Blouin Avocats", HEADING_FONT, Pt(22), NAVY, bold=True)
    p2 = tf.add_paragraph()
    add_run(p2, "plus gros cabinet d'immigration au Quebec", BODY_FONT, Pt(12), MUTED)

    add_slide_number(slide, "05")

    # =========================================================================
    # SLIDE 6: Business Model
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Label
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.0), Inches(4), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "BUSINESS MODEL", BODY_FONT, Pt(11), LIGHT_MUTED, bold=True)

    # Title (two lines)
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.5), Inches(10), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "Tarification a l'usage.", HEADING_FONT, Pt(36), NAVY)
    p2 = tf.add_paragraph()
    add_run(p2, "Qui ", HEADING_FONT, Pt(36), NAVY)
    add_run(p2, "scale", HEADING_FONT, Pt(36), BLUE, italic=True)
    add_run(p2, " avec le cabinet.", HEADING_FONT, Pt(36), NAVY)

    # Hero stat
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(3.2), Inches(8), Inches(2.0))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "15-20", HEADING_FONT, Pt(120), NAVY, bold=True)
    add_run(p, "$", HEADING_FONT, Pt(120), BLUE, bold=True, italic=True)

    # "par demande traitee"
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(5.2), Inches(5), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "par demande traitee", BODY_FONT, Pt(16), MUTED, bold=False)

    # 3 proof points in a row
    proofs = [
        ("2 250$ - 7 500$", "ACV actuel par cabinet"),
        ("30 jours", "essai gratuit"),
        ("2 mois", "offerts sur paiement annuel"),
    ]
    proof_y = Inches(6.0)
    proof_w = Inches(3.3)
    for i, (stat, label) in enumerate(proofs):
        x = LEFT_MARGIN + proof_w * i
        txBox = add_textbox(slide, x, proof_y, proof_w, Inches(1.0))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        add_run(p, stat, HEADING_FONT, Pt(24), NAVY, bold=True)
        p2 = tf.add_paragraph()
        add_run(p2, label, BODY_FONT, Pt(12), MUTED)

    add_slide_number(slide, "06")

    # =========================================================================
    # SLIDE 7: Market
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Label
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.0), Inches(3), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "MARCHE", BODY_FONT, Pt(11), LIGHT_MUTED, bold=True)

    # Title
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.5), Inches(10), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "Le Quebec, seul marche ", HEADING_FONT, Pt(36), NAVY)
    add_run(p, "francophone", HEADING_FONT, Pt(36), BLUE, italic=True)
    add_run(p, " de l'immigration canadienne.", HEADING_FONT, Pt(36), NAVY)

    # 3 stats
    stats = [
        ("500+", "membres AQAADI"),
        ("12K+", "consultants regules au CICC"),
        ("483K", "residents permanents admis en 2024"),
    ]
    stat_y = Inches(3.8)
    stat_w = Inches(3.3)
    for i, (num, label) in enumerate(stats):
        x = LEFT_MARGIN + stat_w * i
        txBox = add_textbox(slide, x, stat_y, stat_w, Inches(2.0))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        add_run(p, num, HEADING_FONT, Pt(54), NAVY, bold=True)
        p2 = tf.add_paragraph()
        add_run(p2, label, BODY_FONT, Pt(12), MUTED)

    # Sources bottom-left
    txBox = add_textbox(slide, LEFT_MARGIN, SLIDE_HEIGHT - Inches(0.7), Inches(7), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Sources : CICC 2024, IRCC Annual Report 2025, AQAADI.", BODY_FONT, Pt(9), LIGHT_MUTED, italic=True)

    add_slide_number(slide, "07")

    # =========================================================================
    # SLIDE 8: Competition (centered)
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Eyebrow centered
    txBox = add_textbox(slide, Inches(0), Inches(0.8), SLIDE_WIDTH, Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "CONCURRENCE", BODY_FONT, Pt(16), LIGHT_MUTED, bold=True)

    # Headline centered
    txBox = add_textbox(slide, Inches(0), Inches(1.5), SLIDE_WIDTH, Inches(1.2))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "Trois concurrents. ", HEADING_FONT, Pt(48), NAVY)
    add_run(p, "Trois angles morts.", HEADING_FONT, Pt(48), BLUE, italic=True)

    # Competitors line
    txBox = add_textbox(slide, Inches(0), Inches(2.8), SLIDE_WIDTH, Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    add_run(p, "Visto AI  ·  VisaFlo  ·  CaseEasy", HEADING_FONT, Pt(20), LIGHT_MUTED)

    # Three columns with ghost numbers and titles
    angles = [
        ("01", "Marche francophone ignore"),
        ("02", "Mauvais probleme resolu"),
        ("03", "Absents des reseaux sociaux"),
    ]
    col_w = Inches(3.8)
    start_x = Inches(1.0)
    angle_y = Inches(4.0)
    for i, (num, title) in enumerate(angles):
        x = start_x + col_w * i
        # Ghost number
        txBox = add_textbox(slide, x, angle_y, col_w, Inches(1.8))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        add_run(p, num, HEADING_FONT, Pt(80), GHOST)

        # Title
        txBox = add_textbox(slide, x + Inches(0.2), angle_y + Inches(1.5), col_w - Inches(0.4), Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        add_run(p, title, HEADING_FONT, Pt(20), NAVY, bold=True)

    add_slide_number(slide, "08")

    # =========================================================================
    # SLIDE 9: Team
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Label
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(0.8), Inches(3), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "EQUIPE", BODY_FONT, Pt(11), LIGHT_MUTED, bold=True)

    # Title
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.3), Inches(10), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "8 mois a l'interne. 4 mois de validation. ", HEADING_FONT, Pt(34), NAVY)
    add_run(p, "Cofondateur en vue.", HEADING_FONT, Pt(34), BLUE, italic=True)

    # Subtitle
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(2.8), Inches(10), Inches(0.6))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "Je connais le workflow d'un cabinet d'immigration mieux que mes utilisateurs.", BODY_FONT, Pt(13), MUTED)

    # Founder block
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(3.6), Inches(5), Inches(0.9))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Parsa Homayouni", HEADING_FONT, Pt(24), NAVY, bold=True)
    p2 = tf.add_paragraph()
    add_run(p2, "Fondateur", BODY_FONT, Pt(13), MUTED, bold=False)

    # 3 proof points
    proofs = [
        ("8 mois", "in-house dans un cabinet d'immigration a observer le workflow"),
        ("4 mois", "de validation avec de vrais cabinets avant ma premiere ligne de code"),
        ("Cofondateur", "en discussion avec un meilleur ami, complementarite technique et business"),
    ]
    proof_y = Inches(5.0)
    proof_w = Inches(3.3)
    for i, (stat, label) in enumerate(proofs):
        x = LEFT_MARGIN + proof_w * i
        txBox = add_textbox(slide, x, proof_y, proof_w, Inches(2.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        add_run(p, stat, HEADING_FONT, Pt(26), NAVY, bold=True)
        p2 = tf.add_paragraph()
        add_run(p2, label, BODY_FONT, Pt(11), MUTED)

    add_slide_number(slide, "09")

    # =========================================================================
    # SLIDE 10: Vision
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Label
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(0.8), Inches(3), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "VISION", BODY_FONT, Pt(11), LIGHT_MUTED, bold=True)

    # Title
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.3), Inches(10), Inches(1.2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "De 60 cabinets a l'infrastructure ", HEADING_FONT, Pt(36), NAVY)
    add_run(p, "nationale.", HEADING_FONT, Pt(36), BLUE, italic=True)

    # Subtitle
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(2.5), Inches(6), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Quatre etapes. Quatre ans.", BODY_FONT, Pt(16), MUTED, bold=False)

    # Left waypoints
    waypoints = [
        ("2026", "60 cabinets -- Quebec"),
        ("2027", "300-500 cabinets -- Canada"),
        ("2028", "Couverture nationale"),
    ]
    wp_y = Inches(3.5)
    wp_gap = Inches(1.0)
    for i, (year, desc) in enumerate(waypoints):
        y = wp_y + wp_gap * i
        txBox = add_textbox(slide, LEFT_MARGIN, y, Inches(4.5), Inches(0.9))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        add_run(p, year, HEADING_FONT, Pt(24), NAVY, bold=True)
        p2 = tf.add_paragraph()
        add_run(p2, desc, BODY_FONT, Pt(13), MUTED)

    # Right climax
    climax_x = Inches(7.0)
    # 2029+ in blue italic
    txBox = add_textbox(slide, climax_x, Inches(3.3), Inches(5), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "2029+", HEADING_FONT, Pt(60), BLUE, bold=True, italic=True)

    # Contrat IRCC
    txBox = add_textbox(slide, climax_x, Inches(4.8), Inches(5), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Contrat IRCC", HEADING_FONT, Pt(24), NAVY, bold=True)

    # Detail
    txBox = add_textbox(slide, climax_x, Inches(5.4), Inches(5), Inches(1.2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "Formio", BRAND_FONT, Pt(12), NAVY, bold=True)
    add_run(p, " devient le handler officiel des demandes d'immigration aupres du gouvernement federal.", BODY_FONT, Pt(12), MUTED)

    add_slide_number(slide, "10")

    # =========================================================================
    # SLIDE 11: Closing / Ask
    # =========================================================================
    slide = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide, BG_COLOR)

    # Wordmark "Formio"
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(1.5), Inches(5), Inches(1.0))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "Formio", BRAND_FONT, Pt(44), NAVY, bold=True)

    # Tagline
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(2.5), Inches(9), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_run(p, "Le futur de l'immigration", HEADING_FONT, Pt(32), NAVY)
    p2 = tf.add_paragraph()
    add_run(p2, "au Canada ", HEADING_FONT, Pt(32), NAVY)
    add_run(p2, "commence au Quebec.", HEADING_FONT, Pt(32), BLUE, italic=True)

    # Ask section label
    txBox = add_textbox(slide, LEFT_MARGIN, Inches(4.5), Inches(4), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "CE QUE JE CHERCHE", BODY_FONT, Pt(11), LIGHT_MUTED, bold=True)

    # 3 ask blocks
    asks = [
        ("Intros", "cabinets et consultants partout au Canada"),
        ("Conseils", "expansion interprovinciale et relations IRCC"),
        ("Partenaires strategiques", "gouvernement federal, associations professionnelles"),
    ]
    ask_y = Inches(5.1)
    ask_w = Inches(3.3)
    for i, (title, desc) in enumerate(asks):
        x = LEFT_MARGIN + ask_w * i
        txBox = add_textbox(slide, x, ask_y, ask_w, Inches(1.2))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        add_run(p, title, HEADING_FONT, Pt(18), NAVY, bold=True)
        p2 = tf.add_paragraph()
        add_run(p2, desc, BODY_FONT, Pt(11), MUTED)

    # Footer URL bottom-left
    txBox = add_textbox(slide, LEFT_MARGIN, SLIDE_HEIGHT - Inches(0.7), Inches(5), Inches(0.4))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    add_run(p, "parsa@formio.ca  ·  formio.ca", BRAND_FONT, Pt(12), NAVY, bold=True)

    add_slide_number(slide, "11")

    # Save
    output_path = "/Users/parsahomayouni/Documents/ProjectsProg/FormioProjects/FormioLandingPage/formio-pitch-deck.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    create_presentation()
